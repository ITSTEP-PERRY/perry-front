"""Build Perry defense PPTX: diagram + description on each slide."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(r"D:\Perry\project_defense")
VIEWS = ROOT / "views_project"
OUT = ROOT / "Perry-Defense-Presentation.pptx"

BG = RGBColor(0x0F, 0x14, 0x19)
FG = RGBColor(0xE8, 0xEE, 0xF5)
MUTED = RGBColor(0xA8, 0xB4, 0xC0)
ACCENT = RGBColor(0x5B, 0x9F, 0xD4)


def set_slide_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_title_bar(slide, text, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.5))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = FG
    run.font.name = "Calibri"
    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.4), Inches(0.65), Inches(12.5), Inches(0.35))
        p2 = box2.text_frame.paragraphs[0]
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(13)
        r2.font.color.rgb = MUTED
        r2.font.name = "Calibri"


def add_bullets(slide, lines, left, top, width, height, size=15, title=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    start = 0
    if title:
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = title
        run.font.size = Pt(size + 1)
        run.font.bold = True
        run.font.color.rgb = ACCENT
        run.font.name = "Calibri"
        p.space_after = Pt(8)
        start = 1
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if (i == 0 and start == 0) else tf.add_paragraph()
        p.text = ("• " if not line.startswith("•") else "") + line.lstrip("• ").strip()
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = FG
        p.font.name = "Calibri"
        p.space_after = Pt(6)


def add_picture_left(slide, path: Path, top=1.1, max_h=5.8, max_w=7.6, left=0.35):
    if not path.exists():
        add_bullets(slide, [f"нет файла: {path.name}"], left, top, max_w, 2, size=14)
        return
    pic = slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(max_w))
    max_h_emu = Inches(max_h)
    if pic.height > max_h_emu:
        ratio = float(max_h_emu) / float(pic.height)
        pic.height = int(pic.height * ratio)
        pic.width = int(pic.width * ratio)


def add_notes(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text


def diagram_slide(prs, blank, title, subtitle, png_path, bullets, notes):
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_title_bar(s, title, subtitle)
    add_picture_left(s, png_path, top=1.1, max_h=5.9, max_w=7.5)
    add_bullets(
        s,
        bullets,
        left=8.1,
        top=1.15,
        width=4.8,
        height=5.9,
        size=14,
        title="Что на схеме",
    )
    add_notes(s, notes)
    return s


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Title
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_title_bar(s, "Perry — интернет-магазин", "Дипломный проект · защита · ITSTEP-PERRY")
    add_bullets(
        s,
        [
            "React-витрина + встроенная админка (:3000)",
            "Product API .NET 8 (:5272) + PostgreSQL",
            "Auth Service и Admin Service — Azure (команда)",
            "Схемы: project_defense/ · демо localhost:3000",
        ],
        left=0.8,
        top=2.0,
        width=11,
        height=4,
        size=22,
    )
    add_notes(s, "Введение: Perry — магазин. Я — Product API и React. Auth/Users — сервисы команды.")

    diagram_slide(
        prs,
        blank,
        "1. Архитектура системы",
        "От браузера до базы данных",
        VIEWS / "01-diagram.png",
        [
            "Browser — пользователь на :3000",
            "Vite + React — SPA витрины и админки",
            "Auth Service (Azure) — login и JWT",
            "Perry.Api :5272 — каталог, корзина, заказы",
            "Admin Service — пользователи админки",
            "PostgreSQL — Product DB",
            "Пунктир Internal #97 — Product → Auth",
        ],
        "Архитектура: React :3000, Product :5272, Auth/Admin Azure, данные в PostgreSQL.",
    )

    diagram_slide(
        prs,
        blank,
        "2. Маршрутизация запросов",
        "Vite proxy: куда уходит каждый путь",
        VIEWS / "02-diagram.png",
        [
            "/ → отдаёт React (HTML/JS/CSS)",
            "/api/* → Product API (JSON)",
            "/auth-api/* → Auth Azure (login/me)",
            "/users-api/* → Admin Service (users)",
            "/uploads/* → файлы картинок на Api",
            "Зачем proxy: same-origin, без CORS в dev",
            "Только /api идёт в PostgreSQL через Api",
        ],
        "Vite проксирует same-origin, чтобы не ловить CORS к Azure Auth.",
    )

    diagram_slide(
        prs,
        blank,
        "3. Стык микросервисов",
        "Карточки #94–#97: Auth · Product · Admin",
        VIEWS / "13-diagram.png",
        [
            "1. Front логинится в Auth → JWT",
            "2. Front бьёт в Product с Bearer",
            "3. Admin Users — отдельный сервис",
            "Product DB без таблицы Users (#94)",
            "UserId берём из claim sub JWT",
            "4. Internal #97 — Product читает имена",
            "Общий HS256 secret у Auth и Product",
        ],
        "Users нет в Product DB. Internal Auth для обогащения заказов.",
    )

    diagram_slide(
        prs,
        blank,
        "4. Backend: структура solution",
        "Perry.sln — слои ответственности",
        VIEWS / "03-diagram.png",
        [
            "Perry.Api — Controllers, JWT, Swagger",
            "Infrastructure — EF, сервисы, storage",
            "Domain — сущности без зависимости от EF",
            "Api вызывает Infrastructure",
            "Infrastructure использует Domain",
            "Perry.Web — legacy Razor (не основной UI)",
            "Основной клиент — React :3000",
        ],
        "Solution .NET 8. Legacy Web не основной UI.",
    )

    diagram_slide(
        prs,
        blank,
        "5. Слой базы данных",
        "EF Core + Npgsql → PostgreSQL",
        VIEWS / "04-diagram.png",
        [
            "Services — Cart, Order, Product…",
            "AppDbContext + Fluent Configurations",
            "EF Migrations — схема таблиц",
            "PostgreSQL 16 (Docker или Supabase)",
            "Сущности: Product, Order, Cart, Review…",
            "Нет Users — только Guid UserId из JWT",
            "Сидер DbSeeder при старте Api",
        ],
        "Миграции PostgreSQL. UserId только из JWT.",
    )

    diagram_slide(
        prs,
        blank,
        "6. Пайплайн HTTP-запроса",
        "Один запрос сверху вниз",
        VIEWS / "06-diagram.png",
        [
            "1. CORS — разрешён origin :3000",
            "2. JWT Bearer — проверка подписи HS256",
            "3. Authorize — нужна ли роль Admin",
            "4. Controller → Service (бизнес-логика)",
            "5. EF пишет/читает PostgreSQL",
            "6. JSON-ответ клиенту",
            "401/403 — если токен или роль не ок",
        ],
        "Стрелки по номерам 1→6 — путь запроса.",
    )

    diagram_slide(
        prs,
        blank,
        "7. Три потока Auth",
        "Покупатель · DEV Admin · Internal",
        VIEWS / "07-diagram.png",
        [
            "A: login на Auth → JWT → Product проверяет",
            "B: DEV Admin/Admin → /api/dev/admin-login",
            "B только в Development, не для прода",
            "C: Product → /internal/auth/token (#97)",
            "C: затем GET /internal/users/{id}",
            "Issuer Perry.AuthService, Audience Perry.Client",
            "Один SigningSecret у Auth и Product",
        ],
        "Тот же JWT, что выдаёт Auth. DEV: Admin/Admin.",
    )

    diagram_slide(
        prs,
        blank,
        "8. Frontend: структура src/",
        "React + Vite без Redux",
        VIEWS / "08-diagram.png",
        [
            "main/App — провайдеры и роутер",
            "app/ — Auth, Cart, Wishlist Context",
            "pages/ — витрина, auth, account, admin",
            "widgets/ — AppShell, карточки, модалки",
            "api/ — client + facade к бэкендам",
            "Состояние — Context, не Redux/Zustand",
            "Стили: storefront.css / admin.css",
        ],
        "Один SPA, состояние через Context.",
    )

    diagram_slide(
        prs,
        blank,
        "9. Маршруты React",
        "Три зоны: витрина · auth · admin",
        VIEWS / "09-diagram.png",
        [
            "Витрина: /, /products, /cart",
            "/account/* — только после login",
            "Auth: /login, /register, forgot…",
            "Admin: /admin/login + CRUD-экраны",
            "AdminShell пускает только isAdmin",
            "Users-страница бьёт в Admin Service",
            "Остальное Product API через /api",
        ],
        "Один SPA: витрина, кабинет, админка.",
    )

    diagram_slide(
        prs,
        blank,
        "10. Корзина → заказ",
        "Главный пользовательский сценарий",
        VIEWS / "12-diagram.png",
        [
            "1. Гостю выдаём sessionId (UUID)",
            "2. Add to cart → POST /api/cart",
            "3. После login — merge корзины",
            "4. Checkout → POST /api/orders/checkout",
            "5. Редирект в /account/orders",
            "Заказ пишется в PostgreSQL",
            "Admin видит заказы + lastUpdateUtc",
        ],
        "Гостевая корзина → merge → checkout → заказы.",
    )

    diagram_slide(
        prs,
        blank,
        "11. Что сделали в спринте",
        "Ключевые доработки для защиты",
        VIEWS / "14-diagram.png",
        [
            "#A08 — Product API на PostgreSQL",
            "#A09/#A10 — admin orders + last update",
            "Admin Categories — UI под Figma",
            "CI — compose validate + gitleaks",
            "Схемы и речь — project_defense/",
            "Trello #92 — графические слайды",
            "Код в perry-front и My_Amazon2",
        ],
        "PG, orders, categories, CI.",
    )

    # Demo
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_title_bar(s, "12. Демо на защите", "localhost:3000 · Swagger :5272")
    add_bullets(
        s,
        [
            "Главная / каталог товаров",
            "Корзина → checkout (или login)",
            "Админка: Admin / Admin → категории / заказы",
            "Swagger: /api/health, /api/products",
            "Честно: SMTP stub; Auth Azure — сервис команды",
        ],
        left=0.8,
        top=1.8,
        width=11,
        height=4.5,
        size=20,
    )
    add_notes(s, "Демо 1–2 мин.")

    # Q&A
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_title_bar(s, "Спасибо. Вопросы?", "JWT · Internal #97 · PostgreSQL · корзина")
    add_bullets(
        s,
        [
            "Материалы: project_defense/ + DEFENSE_QA.md",
            "Front: ITSTEP-PERRY/perry-front",
            "API: Teslyar75/My_Amazon2 · ITSTEP Back_end",
            "Trello #92 — слайды по схемам",
        ],
        left=0.8,
        top=2.0,
        width=11,
        height=4,
        size=20,
    )
    add_notes(s, "Заключение 20 сек.")

    prs.save(str(OUT))
    print(f"Wrote {OUT} ({OUT.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
